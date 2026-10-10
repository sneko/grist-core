"""
Tests direct_sandbox, which runs the data engine in-process (as browsers do for encrypted
documents, in Pyodide).
"""
import collections
import marshal
import types
import unittest

import direct_sandbox
import sandbox as sandbox_module
from sandbox import Sandbox


# Creating a direct sandbox makes it the default one: restore the default for other tests.
_default_sandbox = []

def setUpModule():
  _default_sandbox.append(sandbox_module.default_sandbox)

def tearDownModule():
  sandbox_module.default_sandbox = _default_sandbox.pop()


class Engine(object):
  """A copy of the engine, driven the way browsers drive it."""
  def __init__(self):
    self.sandbox = direct_sandbox.create()

  def call(self, *args):
    code, data = marshal.loads(self.sandbox.call(marshal.dumps(list(args), 2)))
    if code != Sandbox.DATA:
      raise Exception(data)
    return data

  def stored(self, result):
    return [action for (_env, action) in result['stored']]

  @property
  def engine(self):
    # pylint: disable=protected-access
    return self.sandbox._functions['load_empty'].__wrapped__.__self__

  def values(self, table_id, col_id):
    return self.call('fetch_table', table_id)[3][col_id]


def new_doc():
  eng = Engine()
  eng.call('load_empty')
  eng.call('apply_user_actions', [['InitNewDoc']])
  return eng


def load_doc(tables):
  eng = Engine()
  others = eng.call('load_meta_tables_repr', tables['_grist_Tables'], tables['_grist_Tables_column'])
  for table_id in others:
    eng.call('load_table_repr', tables.get(table_id, ['TableData', table_id, [], {}]))
  eng.call('calculate_exactly')
  return eng


class TestDirectSandbox(unittest.TestCase):
  def test_calls(self):
    eng = new_doc()
    self.assertEqual(eng.call('test_echo', 'hello'), 'hello')
    with self.assertRaisesRegex(Exception, 'KeyError'):
      eng.call('no_such_function')
    with self.assertRaisesRegex(Exception, 'KeyError'):
      eng.call('apply_user_actions', [['AddRecord', 'NoSuchTable', None, {}]])

  def test_reload_and_replay_exactly(self):
    eng = new_doc()
    eng.call('apply_user_actions', [
      ['AddTable', 'T', [
        {'id': 'A', 'type': 'Text', 'isFormula': False},
        {'id': 'Upper', 'type': 'Any', 'isFormula': True, 'formula': '$A.upper()'},
      ]],
      ['BulkAddRecord', 'T', [None, None], {'A': ['x', 'y']}],
    ])
    tables = eng.call('fetch_all_tables')
    self.assertIn('T', tables)
    self.assertIn('_grist_DocInfo', tables)

    # Another copy loads the document from its tables, and ends up the same.
    eng2 = load_doc(tables)
    self.assertEqual(eng2.call('fetch_all_tables'), tables)

    # A change with random results, replayed by the other copy, gives the same values there.
    change = eng.call('apply_user_actions', [
      ['AddColumn', 'T', 'Rand', {'type': 'Any', 'isFormula': True, 'formula': 'RANDBETWEEN(1, 1000000000)'}],
      ['AddRecord', 'T', None, {'A': 'z'}],
    ])
    eng2.call('apply_doc_actions_exactly', eng.stored(change))
    rand = eng.call('fetch_table', 'T')[3]['Rand']
    self.assertTrue(all(isinstance(v, int) for v in rand))
    self.assertEqual(eng2.call('fetch_table', 'T'), eng.call('fetch_table', 'T'))

    # Formulas added by a replayed change keep track of what they depend on.
    added = eng.call('apply_user_actions', [
      ['AddColumn', 'T', 'Init', {'type': 'Any', 'isFormula': True, 'formula': '$A[:1]'}]])
    eng2.call('apply_doc_actions_exactly', eng.stored(added))
    result = eng2.call('apply_user_actions', [['UpdateRecord', 'T', 1, {'A': 'w'}]])
    self.assertIn(['UpdateRecord', 'T', 1, {'Init': 'w'}], eng2.stored(result))

  def test_load_keeps_values(self):
    eng = new_doc()
    eng.call('apply_user_actions', [
      ['AddTable', 'T', [
        {'id': 'A', 'type': 'Text', 'isFormula': False},
        {'id': 'Rand', 'type': 'Any', 'isFormula': True, 'formula': 'RANDBETWEEN(1, 1000000000)'},
        {'id': 'Upper', 'type': 'Any', 'isFormula': True, 'formula': '$A.upper()'},
      ]],
      ['BulkAddRecord', 'T', [None, None], {'A': ['x', 'y']}],
    ])
    tables = eng.call('fetch_all_tables')
    # Another copy loads the document: it keeps the values of formulas, even random ones...
    eng2 = load_doc(tables)
    self.assertEqual(eng2.call('fetch_table', 'T'), eng.call('fetch_table', 'T'))
    # ... while knowing what formulas depend on.
    result = eng2.call('apply_user_actions', [['UpdateRecord', 'T', 1, {'A': 'w'}]])
    self.assertIn(['UpdateRecord', 'T', 1, {'Upper': 'W'}], eng2.stored(result))
    self.assertEqual(eng2.call('fetch_table', 'T')[3]['Rand'], tables['T'][3]['Rand'])

  def test_frozen_columns(self):
    eng = new_doc()
    eng.call('apply_user_actions', [
      ['AddTable', 'T', [
        {'id': 'A', 'type': 'Text', 'isFormula': False},
        {'id': 'F', 'type': 'Any', 'isFormula': True, 'formula': '$A.upper()'},
      ]],
      ['AddRecord', 'T', None, {'A': 'x'}],
    ])
    eng.call('set_frozen_columns', [['T', 'F']])
    eng.call('apply_user_actions', [['UpdateRecord', 'T', 1, {'A': 'y'}]])
    # Frozen formulas keep the values they're given, rather than being recomputed.
    self.assertEqual(eng.call('fetch_table', 'T')[3]['F'], ['X'])
    eng.call('apply_doc_actions_exactly', [['UpdateRecord', 'T', 1, {'F': 'given'}]])
    self.assertEqual(eng.call('fetch_table', 'T')[3]['F'], ['given'])
    eng.call('set_frozen_columns', [])
    eng.call('apply_user_actions', [['UpdateRecord', 'T', 1, {'A': 'z'}]])
    self.assertEqual(eng.call('fetch_table', 'T')[3]['F'], ['Z'])


def error_message(eng, col_id, row_id=1):
  # Formula errors are encoded as ['E', exception name, message, ...].
  error = eng.call('get_formula_error', 'Staff', col_id, row_id)
  return "%s: %s" % (error[1], error[2]) if isinstance(error, list) and error[:1] == ['E'] else None


def staff_doc():
  eng = new_doc()
  eng.call('apply_user_actions', [
    ['AddTable', 'Staff', [
      {'id': 'Name', 'type': 'Text', 'isFormula': False},
      {'id': 'Salary', 'type': 'Numeric', 'isFormula': False},
    ]],
    ['BulkAddRecord', 'Staff', [None, None], {'Name': ['Ana', 'Ben'], 'Salary': [52100, 61300]}],
  ])
  return eng


def add_formula(eng, col_id, formula):
  eng.call('apply_user_actions', [
    ['AddColumn', 'Staff', col_id, {'type': 'Any', 'isFormula': True, 'formula': formula}]])


class TestFormulaLockdown(unittest.TestCase):
  def test_regular_formulas(self):
    eng = staff_doc()
    add_formula(eng, 'Upper', '$Name.upper()')
    add_formula(eng, 'Label', '"{} earns {:.0f}".format($Name, $Salary)')
    add_formula(eng, 'Fstr', 'f"{$Name}: {$Salary:.0f}"')
    add_formula(eng, 'Found', 'Staff.lookupOne(Name="Ben").Salary')
    add_formula(eng, 'Re', 're.sub("a", "o", $Name)')
    add_formula(eng, 'Date', 'datetime.date(2024, 1, 2).year + math.floor(1.5)')
    self.assertEqual(eng.values('Staff', 'Upper'), ['ANA', 'BEN'])
    self.assertEqual(eng.values('Staff', 'Label'), ['Ana earns 52100', 'Ben earns 61300'])
    self.assertEqual(eng.values('Staff', 'Fstr'), ['Ana: 52100', 'Ben: 61300'])
    self.assertEqual(eng.values('Staff', 'Found'), [61300, 61300])
    self.assertEqual(eng.values('Staff', 'Re'), ['Ano', 'Ben'])
    self.assertEqual(eng.values('Staff', 'Date'), [2025, 2025])

  def test_introspection_is_refused(self):
    eng = staff_doc()
    formulas = {
      'A': 'getattr(rec, "Sal" + "ary")',
      'B': 'rec._table',
      'C': 'import gc\nreturn gc.get_objects()',
      'D': '"{0._table}".format(rec)',
      'E': 'fmt = "{0.Name}"\nreturn fmt.format(rec)',
      'F': 'vars(rec)',
      'G': '(x for x in [1]).gi_frame',
      'H': 'rec.__class__',
      'K': 'class K:\n  def __getattr__(self, n): return n\nreturn K().x',
      'L': 'f = lambda __x: __x\nreturn f(1)',
    }
    for col_id, formula in formulas.items():
      add_formula(eng, col_id, formula)
      self.assertRegex(error_message(eng, col_id) or '', r'^PermissionError: In encrypted documents',
                       formula)
    # Going around the engine fails at runtime.
    add_formula(eng, 'I', 'Staff.table.get_column("Sal" + "ary").raw_get(1)')
    self.assertRegex(error_message(eng, 'I'), r'^AttributeError')
    add_formula(eng, 'J', 're.enum')
    self.assertRegex(error_message(eng, 'J'), r'^AttributeError')

  def test_nothing_internal_is_reachable(self):
    # Explore what formulas can reach (by names they may use), looking for the engine's internals.
    eng = staff_doc()
    with direct_sandbox.formula_lockdown():
      engine = eng.engine
      usercode = engine.gencode.usercode
      staff = usercode.Staff
      roots = {name: value for name, value in vars(usercode).items() if not name.startswith('_')}
      roots.update({name: value for name, value in vars(__builtins__ if isinstance(__builtins__, dict)
                                                     else vars(__builtins__)).items()
                    if not name.startswith('_')} if False else {})
      import builtins
      roots.update({name: getattr(builtins, name) for name in dir(builtins) if not name.startswith('_')})
      roots['rec'] = staff.Record(1)
      roots['records'] = staff.all
    internal = ('Engine', 'Table', 'BaseColumn', 'DocActions', 'UserActions', 'DocModel', 'frame')
    seen = {}
    found = []
    queue = collections.deque((value, name) for name, value in roots.items()
                              if name not in direct_sandbox.DENIED_NAMES)
    while queue:
      obj, path = queue.popleft()
      if id(obj) in seen or path.count('.') > 4:
        continue
      seen[id(obj)] = obj
      if any(c.__name__ in internal for c in type(obj).__mro__) or isinstance(obj, types.ModuleType):
        found.append(path)
        continue
      if isinstance(obj, (str, bytes, int, float)):
        continue
      names = set()
      for c in type(obj).__mro__:
        names.update(vars(c))
      try:
        names.update(vars(obj))
      except TypeError:
        pass
      if not isinstance(obj, type):
        names.update(dir(obj))
      for name in sorted(names):
        if name.startswith('_') or name in direct_sandbox.DENIED_NAMES:
          continue
        if name == 'table' and isinstance(obj, type(staff)):
          continue    # Refused at runtime (see test_introspection_is_refused).
        try:
          value = getattr(obj, name)
        except Exception:  # pylint: disable=broad-except
          continue
        queue.append((value, path + '.' + name))
    self.assertEqual(found, [])


class TestExportXlsx(unittest.TestCase):
  def test_export_xlsx(self):
    import base64, io  # pylint: disable=import-outside-toplevel,multiple-imports
    import openpyxl  # pylint: disable=import-outside-toplevel
    eng = new_doc()
    data = eng.call('export_xlsx', [
      ['Staff', ['Name', 'Salary', 'Active'], [['Ana', 52100.5, True], ['=1+1', None, False]]],
      ['Staff', ['A'], [['x']]],
    ])
    workbook = openpyxl.load_workbook(io.BytesIO(base64.b64decode(data)))
    self.assertEqual(workbook.sheetnames, ['Staff', 'Staff (2)'])
    rows = [[c.value for c in row] for row in workbook['Staff'].iter_rows()]
    self.assertEqual(rows, [['Name', 'Salary', 'Active'], ['Ana', 52100.5, True], ['=1+1', None, False]])
    # Text that looks like a formula stays text.
    self.assertEqual(workbook['Staff']['A3'].data_type, 's')


class TestReservedRowIds(unittest.TestCase):
  def test_removed_private_column_ids_are_not_reused(self):
    eng = staff_doc()
    columns = lambda: eng.call('fetch_table', '_grist_Tables_column')
    salary_ref = columns()[2][columns()[3]['colId'].index('Salary')]
    # The browser says Salary's id must not be reused (it's a private column's).
    eng.call('set_reserved_row_ids', '_grist_Tables_column', [salary_ref])
    removal = eng.call('apply_user_actions', [['RemoveColumn', 'Staff', 'Salary']])
    addition = eng.call('apply_user_actions', [['AddColumn', 'Staff', 'Other', {'type': 'Text'}]])
    other_ref = columns()[2][columns()[3]['colId'].index('Other')]
    self.assertGreater(other_ref, salary_ref)
    # Undoing the removal restores the column with its id.
    undo = lambda result: eng.call('apply_user_actions', [['ApplyUndoActions', [a for (_e, a) in result['undo']]]])
    undo(addition)
    undo(removal)
    self.assertEqual(columns()[2][columns()[3]['colId'].index('Salary')], salary_ref)


class TestScopeMonitor(unittest.TestCase):
  def test_formulas_only_read_private_columns_they_mention(self):
    eng = staff_doc()
    add_formula(eng, 'Double', '$Salary * 2')
    add_formula(eng, 'Sneaky', 'Staff.lookupOne(**{"Sal" + "ary": 61300}).Name')
    add_formula(eng, 'Sorted', 'Staff.lookupRecords(order_by="-Sal" + "ary").Name')
    add_formula(eng, 'Public', 'Staff.lookupOne(Name=$Name).id')
    # The browser tells the engine which columns are private (here, by mention: Double).
    eng.call('set_column_scopes', [['Staff', 'Salary', ['column:x']], ['Staff', 'Double', ['column:x']]])
    eng.call('apply_user_actions', [['UpdateRecord', 'Staff', 1, {'Salary': 52200}, ]])
    eng.call('apply_user_actions', [['BulkUpdateRecord', 'Staff', [1, 2], {'Name': ['Ann', 'Bob']}]])
    self.assertEqual(eng.values('Staff', 'Double'), [104400, 122600])
    self.assertEqual(eng.values('Staff', 'Public'), [1, 2])
    for col_id in ('Sneaky', 'Sorted'):
      self.assertRegex(error_message(eng, col_id) or '',
                       r'^PermissionError: This formula reads a private column without naming it')
    # Once the column isn't private, they work again.
    eng.call('set_column_scopes', [])
    eng.call('apply_user_actions', [['BulkUpdateRecord', 'Staff', [1, 2], {'Name': ['Ana', 'Ben']}]])
    self.assertEqual(eng.values('Staff', 'Sneaky'), ['Ben', 'Ben'])


if __name__ == "__main__":
  unittest.main()
